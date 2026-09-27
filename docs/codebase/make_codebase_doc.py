# -*- coding: utf-8 -*-
"""
make_codebase_doc.py — the codebase, explained, as a .docx and a .pdf.

    python docs/codebase/make_codebase_doc.py

Most of this document is READ FROM THE CODE rather than typed: the module
table, the line counts, the public API, the network shape and every training
constant come from importing the package and walking its source.  So the
document cannot quietly drift from what the code does — re-run it after a
change and the numbers move with it.

The prose between those generated parts is the part a person wrote: what each
stage is for and which guarantee it enforces.

Windows only for the PDF step: it drives the installed Word, the same way
paper_figures/programs/make_pdfs.py does.
"""
import ast
import io
import os
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
PKG = os.path.join(ROOT, "src", "csdrecon")
SCRIPTS = os.path.join(ROOT, "scripts")
FIGURES = os.path.join(ROOT, "docs", "figures")
OUT_DOCX = os.path.join(HERE, "codebase.docx")

sys.path.insert(0, os.path.join(ROOT, "src"))

BODY_PT = 9.5
FONT = "Calibri"
MONO = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x60, 0x60, 0x60)


# ── reading the code ──────────────────────────────────────────────────────

def modules():
    """[(import path, file path, summary, lines, [public names])], sorted."""
    out = []
    for dirpath, _, names in os.walk(PKG):
        if "__pycache__" in dirpath:
            continue
        for name in sorted(names):
            if not name.endswith(".py"):
                continue
            path = os.path.join(dirpath, name)
            rel = os.path.relpath(path, os.path.dirname(PKG)).replace("\\", "/")
            src = io.open(path, encoding="utf-8").read()
            tree = ast.parse(src)
            doc = ast.get_docstring(tree) or ""
            summary = _summary(doc)
            public = [n.name for n in tree.body
                      if isinstance(n, (ast.FunctionDef, ast.ClassDef))
                      and not n.name.startswith("_")]
            out.append((rel, path, summary, len(src.splitlines()), public))
    return sorted(out)


def _summary(doc, strip_name=True):
    """The opening sentence of a docstring, as one line.

    Docstrings here wrap, so the sentence is taken from the joined first
    paragraph rather than from its first line — otherwise a summary comes out
    cut in half.
    """
    if not doc:
        return ""
    lines = []
    for raw in doc.splitlines():
        if raw.strip():
            lines.append(raw.strip())
        elif lines:
            break
    text = " ".join(lines)
    if strip_name:
        head = text.split(". ")[0]
        for dash in ("—", "--"):
            if dash in head:
                text = text.split(dash, 1)[1].strip()
                break
    stop = text.find(". ")
    if stop > 0:
        text = text[:stop + 1]
    return text[:1].upper() + text[1:] if text else ""


def run_scripts():
    """[(file, first line of its docstring)] for scripts/run_*.py, in order."""
    out = []
    for name in sorted(os.listdir(SCRIPTS)):
        if not (name.startswith("run_") and name.endswith(".py")):
            continue
        doc = ast.get_docstring(ast.parse(
            io.open(os.path.join(SCRIPTS, name), encoding="utf-8").read())) or ""
        out.append((name, _summary(doc, strip_name=False)))
    return sorted(out, key=lambda t: _run_order(t[0]))


def _run_order(name):
    part = name[len("run_"):].split("_")[0]
    return (int(part) if part.isdigit() else 99, name)


def live_facts():
    """The network and the training constants, from the code itself."""
    import torch  # noqa: F401  (imported by grid_model)
    from csdrecon.ml import grid_model, grid_train
    from csdrecon.ml.ray_peaks import NET_CHANNELS
    net = grid_model.RayToLinesNet(in_channels=NET_CHANNELS, depth=3)
    return {
        "widths": net.widths,
        "depth": net.depth,
        "in_channels": net.in_channels,
        "params": sum(p.numel() for p in net.parameters()),
        "width_base": grid_model.WIDTH,
        "train": grid_train.training_description(),
        "thresholds": grid_train.THRESHOLDS,
        "val_fraction": grid_train.VAL_FRACTION,
        "max_pos_weight": grid_train.MAX_POS_WEIGHT,
        "batch": grid_train.BATCH_SIZE,
        "lr": grid_train.LEARNING_RATE,
        "seed": grid_train.SEED,
    }


def line_fraction_text():
    """(line pixels as a %, what an all-zero prediction then scores).

    Read from the newest run's comparison.csv, so the two numbers in the
    metric argument are the run's and not a remembered pair.
    """
    import csv
    import glob
    found = sorted(glob.glob(os.path.join(ROOT, "results", "*",
                                          "comparison.csv")),
                   key=os.path.getmtime)
    if not found:
        return "a few percent", "over 0.9"
    rows = list(csv.DictReader(io.open(found[-1], encoding="utf-8")))
    fracs = {float(r["true_line_fraction"]) for r in rows}
    f = sum(fracs) / len(fracs)
    return "%.1f %%" % (100 * f), "%.3f" % (1 - f)


def n_caps():
    """How many capacitance entries a device actually draws.

    Counted from the split's own VECTOR_KEYS rather than typed: the count was
    written as 13 in three places on 2026-09-21 and is 14.
    """
    from csdrecon.study.device_split import VECTOR_KEYS
    return len(VECTOR_KEYS)


def pool_example():
    """A real pool folder name, read off disk rather than typed."""
    pools = os.path.join(ROOT, "data", "_device_pools")
    if os.path.isdir(pools):
        names = sorted(n for n in os.listdir(pools)
                       if os.path.isdir(os.path.join(pools, n)))
        if names:
            return names[0]
    return "devices_n<N>_res<R>_c<fingerprint>"


def test_count():
    """How many test functions the suite has, without running it."""
    n = 0
    tests = os.path.join(ROOT, "tests")
    for name in sorted(os.listdir(tests)):
        if name.startswith("test_") and name.endswith(".py"):
            tree = ast.parse(io.open(os.path.join(tests, name),
                                     encoding="utf-8").read())
            n += sum(1 for x in ast.walk(tree)
                     if isinstance(x, ast.FunctionDef)
                     and x.name.startswith("test_"))
    return n


# ── writing the document ──────────────────────────────────────────────────

def setup(doc):
    s = doc.styles["Normal"]
    s.font.name = FONT
    s.font.size = Pt(BODY_PT)
    s.font.color.rgb = INK
    pf = s.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)


def title(doc, text, subtitle):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(22)
    r.font.bold = True
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(10.5)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(14)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(13.5)
    r.font.bold = True


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.bold = True


def para(doc, text, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    # **bold** spans
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            p.add_run(chunk).font.bold = bool(i % 2)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.22)
    r = p.add_run(text)
    r.font.name = MONO
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x20, 0x35, 0x50)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for cell, head in zip(t.rows[0].cells, headers):
        cell.text = ""
        r = cell.paragraphs[0].add_run(head)
        r.font.bold = True
        r.font.size = Pt(8.5)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(value))
            r.font.size = Pt(8.5)
    if widths:
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(doc, name, caption, width=5.4):
    path = os.path.join(FIGURES, name)
    if not os.path.isfile(path):
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    c.paragraph_format.space_after = Pt(10)
    return True


# ── the document ──────────────────────────────────────────────────────────

def build():
    mods = modules()
    facts = live_facts()
    doc = Document()
    setup(doc)

    total = sum(m[3] for m in mods)
    title(doc, "probabilistic-csd-reconstruction",
          "How the code works: the pipeline, the method it implements, and "
          "where every number comes from.\n"
          "E. Alizadeh Kashtiban, T. Fujita, A. Oiwa — Osaka University "
          "(SANKEN).  Generated from the source by "
          "docs/codebase/make_codebase_doc.py.")

    # 1 ────────────────────────────────────────────────────────────────────
    h1(doc, "1.  What the code does")
    para(doc,
         "The study recovers the charge-transition lines of a simulated "
         "double quantum dot from a few percent of the gate-voltage plane. "
         "The plane is measured along rays, a U-Net turns what the rays saw "
         "into one probability per pixel, and the probability map is cut at "
         "a threshold to give transition lines that can be scored against "
         "the simulator's exact answer.")
    para(doc,
         "Everything hangs on one comparison being fair: two measurement "
         "budgets must differ in the measurement and in nothing else. Most "
         "of the design below exists to enforce that — one device pool "
         "shared by every budget, one split stored with the pool, one "
         "architecture with one parameter count, and one set of training "
         "constants.")
    figure(doc, "unet_io.png",
           "What the network is given and what it returns: two sparse "
           "channels in, one probability per pixel out.")

    para(doc, "The chain, in the order the programs run it:")
    code(doc,
         "capacitances  ->  simulated device  ->  rays  ->  2 channels\n"
         "                                                     |\n"
         "                                                   U-Net\n"
         "                                                     |\n"
         "                                          probability per pixel\n"
         "                                                     |\n"
         "                                    threshold  ->  transition lines\n"
         "                                                     |\n"
         "                                      F1 at tolerance tau, beside\n"
         "                                      the area that tolerance covers")

    # 2 ────────────────────────────────────────────────────────────────────
    h1(doc, "2.  Layout")
    para(doc,
         "**Nothing in src/ has a command line.** The programs are the "
         "run_*.py files in scripts/: each is a settings block followed by a "
         "few calls into the library. That separation is deliberate — a "
         "checkpoint is scored by code that cannot have seen the test data, "
         "and a re-score can never silently retrain.")
    rows = [(name, first) for name, first in run_scripts()]
    table(doc, ["program", "what it does"], rows, widths=[1.8, 4.6])

    para(doc,
         "The library is %d lines across %d modules:" % (total, len(mods)))
    table(doc, ["module", "lines", "what it is"],
          [(m[0].replace("csdrecon/", ""), m[3],
            m[2] or "(package marker — re-exports only)") for m in mods],
          widths=[1.9, 0.5, 4.0])

    # 3 ────────────────────────────────────────────────────────────────────
    h1(doc, "3.  Stage 1 — the devices, the split and the measurement")
    h2(doc, "The device pool")
    para(doc,
         "A device is one uniform draw of %d capacitances from the " % n_caps() +
         "intervals in config/capacitance_config.py, simulated under the "
         "constant-capacitance model. **No draw is filtered, inspected or "
         "redrawn**: there is no acceptance test anywhere in the pipeline, "
         "so the pool is a plain uniform sample of the parameter box and "
         "its acceptance rate is 1.00 by construction. A draw is lost only "
         "if the simulator itself raises.")
    para(doc,
         "The gate-voltage window is randomly offset per device, so the "
         "honeycomb is not registered to the image frame and the network "
         "cannot learn a fixed position for the lines.")
    para(doc,
         "Pools live in data/_device_pools/ and are shared by every budget, "
         "because devices do not depend on how many rays are fired at them. "
         "The folder name carries a fingerprint of the exact capacitance "
         "intervals, so changing the parameter space builds a new pool "
         "instead of silently reusing the old one.")
    figure(doc, "dqd_model.png",
           "The simulated device: the capacitor network and the four "
           "capacitance matrices one draw fills in.", width=5.0)

    h2(doc, "The split, and the leak it prevents")
    para(doc,
         "**The thing that is split is the device, not the image.** "
         "Capacitance configurations are drawn first, each gets an ID, and "
         "the IDs are split once and written into the pool folder. Every "
         "image and every measurement budget inherits the ID of the device "
         "it came from.")
    para(doc,
         "This matters because a sweep reuses one cached pool across every "
         "(rays, points) cell. If each cell re-split, device 37 could train "
         "in the 4-ray cell and be held out in the 8-ray cell, and the "
         "comparison across cells would stop being like for like. So the "
         "split is decided once, outside the sweep, and read back — never "
         "recomputed (study/device_split.py).")
    para(doc,
         "separation() then answers the obvious referee question with a "
         "number rather than an argument: the smallest Euclidean distance "
         "between any training and any test device in normalised parameter "
         "space, reported beside the same statistic computed within the "
         "training set as a yardstick.")

    h2(doc, "The measurement")
    para(doc,
         "A budget is n_rays x n_points. The rays fan out from the "
         "(max Vx, max Vy) corner at equally spaced angles — "
         "linspace(0, 90, R+2)[1:-1] degrees — and each is sampled at "
         "n_points equally spaced positions, assigned to the nearest pixel "
         "centre (ml/ray_peaks.py).")
    para(doc,
         "Because all rays share an origin, points collide near the corner: "
         "the number of **distinct** pixels touched is slightly smaller "
         "than n_rays x n_points, and coverage is always quoted as distinct "
         "pixels.")
    para(doc, "The network is handed two channels, both full size:")
    para(doc, "ch0 — the raw sensor value at every pixel a ray visited, "
              "zero elsewhere", bullet=True)
    para(doc, "ch1 — a binary visited mask, 1 where any ray passed",
         bullet=True)
    para(doc,
         "A third channel holds the peaks a 1-D peak finder found along "
         "each trace. **It is not network input** — it exists for the "
         "measurement figures and for the classical baseline.")
    para(doc,
         "ch1 is what separates “measured here, and the signal was "
         "low” from “never looked here”. Without it a zero in "
         "ch0 is ambiguous at every pixel and the network cannot tell "
         "absence of evidence from evidence of absence.")
    para(doc,
         "The input shape is the same for every budget: the budget changes "
         "how sparse the two channels are, never how big. One architecture "
         "with one parameter count therefore serves the whole sweep, which "
         "is what makes a difference across it a difference in measurement.")

    # 4 ────────────────────────────────────────────────────────────────────
    h1(doc, "4.  Stage 2 — the network and the training")
    t = facts["train"]
    para(doc,
         "RayToLinesNet (ml/grid_model.py) is a U-Net of depth %d, "
         "channel widths %s, %s trainable parameters. It is fully "
         "convolutional — no flatten, no fixed-size layer — so any grid "
         "size works and every budget uses the identical network. "
         "Upsampling resizes to the skip tensor's own size rather than by a "
         "fixed factor, which is what lets an odd ladder "
         "(100 -> 50 -> 25 -> 12) reassemble exactly."
         % (facts["depth"], "/".join(str(w) for w in facts["widths"]),
            format(facts["params"], ",")))
    table(doc, ["setting", "value", "why it is that"], [
        ("loss", t["loss"],
         "line pixels are a few percent of the diagram, so unweighted BCE "
         "converges to “no lines anywhere” and never leaves"),
        ("max_pos_weight", facts["max_pos_weight"],
         "the raw ~1:30 imbalance makes “call everything a "
         "line” the cheapest early minimum, where the network stops "
         "depending on its input at all"),
        ("optimiser", "%s, lr %g" % (t["optimizer"], facts["lr"]), ""),
        ("batch size", facts["batch"], ""),
        ("val_fraction", facts["val_fraction"],
         "carved out of the TRAINING devices, never the test set"),
        ("threshold scan", ", ".join(str(x) for x in facts["thresholds"]),
         "scanned on the validation split; the best is stored in the "
         "checkpoint"),
        ("checkpoint kept", "best validation F1@1",
         "not the last epoch — a run can end on a high loss and still have "
         "saved a good model"),
    ], widths=[1.25, 1.75, 3.4])
    para(doc,
         "**The test set never enters stage 2.** The validation slice that "
         "picks both the epoch and the binarisation threshold is carved out "
         "of the training devices, which is what makes stage 3's number an "
         "honest held-out result.")

    # 5 ────────────────────────────────────────────────────────────────────
    h1(doc, "5.  Stage 3 — threshold, tolerance and the score")
    h2(doc, "Why not 0.5, and why not here")
    para(doc,
         "The last layer is a sigmoid, so a prediction is a probability "
         "map; every reported number is computed on a binary map cut from "
         "it. A fixed 0.5 would be wrong: the loss weights the positive "
         "class, which deliberately pushes probabilities up, and the best "
         "operating point moves with how sparse the measurement is. The cut "
         "is a fitted quantity, so it is chosen out of sample and stored in "
         "the checkpoint before any test device is seen.")
    figure(doc, "probability.png",
           "What the network outputs before thresholding, and where the "
           "chosen cut sits.")

    h2(doc, "Why the score is tolerant")
    para(doc,
         "Transition lines are one pixel wide and %s of the diagram, "
         "which breaks the obvious metrics. **Pixel accuracy is useless** — "
         "predicting no line anywhere already scores %s — and it is "
         "never quoted as a result. Strict pixel F1 is unfairly harsh: a "
         % line_fraction_text() +
         "perfectly recovered line displaced by one pixel shares no pixel "
         "with the truth and scores zero, although the reconstruction is "
         "physically correct.")
    para(doc,
         "So the headline is tolerant F1 at tau = 1 pixel: a predicted "
         "pixel counts as correct if a true one lies within tau, and a true "
         "pixel counts as found if a predicted one lies within tau. "
         "Distances are Euclidean, so a diagonal neighbour (1.41 px) only "
         "counts from tau = 2. It is implemented with a distance transform "
         "— one EDT per map, not a comparison of every pixel pair.")
    para(doc, "F1 at tau can only rise with tau, so it is never quoted "
              "alone. Two companions are reported with it:")
    para(doc, "**coverage@tau** — the fraction of the plane within tau "
              "pixels of a MEASURED pixel. It describes the experiment and "
              "is the same number whatever the model predicts.", bullet=True)
    para(doc, "**claimed@tau** — the fraction within tau pixels of a "
              "PREDICTED pixel. It depends on the model, and it is the "
              "price of the tolerance: when the discs swallow half the "
              "plane, a high F1@tau says very little.", bullet=True)
    para(doc,
         "Strict IoU is reported too. It punishes one-pixel-wide structures "
         "severely and is included for completeness rather than as a "
         "headline.")

    # 6 ────────────────────────────────────────────────────────────────────
    h1(doc, "6.  Stage 4 — comparing budgets")
    para(doc,
         "comparison.py reads the metrics.json each evaluation wrote and "
         "puts every configuration in one table, in a folder named after "
         "the sweep. The name is a pure function of the sweep, so "
         "re-running one updates its folder and two different sweeps can "
         "never land in the same place.")
    para(doc,
         "Configurations that have not been evaluated are listed as "
         "missing rather than dropped, so a half-finished sweep cannot "
         "quietly look complete. Every row carries the split mode, the "
         "resolution and the training-set size, and the comparison warns "
         "when they are not constant across the rows being plotted — two "
         "budgets are comparable only if everything except the budget "
         "matches.")
    para(doc,
         "A run that never left its initial plateau still produces a full "
         "set of metrics and is indistinguishable in a table from a budget "
         "that is simply too small. It is not: it is a failed fit. "
         "converged() marks those on the validation score, and they are "
         "reported separately rather than dropped — publishing one unmarked "
         "would invite the reader to conclude that more rays made things "
         "worse.")
    para(doc,
         "The gallery is rendered one bundle at a time (figure_bundles.py) "
         "— 40_points_per_ray/, 50_points_per_ray/, and so on. Fifteen "
         "labelled series on one axes exhausts the palette and the labels "
         "collide, so the combined version was unreadable long before it "
         "was wrong.")

    # 7 ────────────────────────────────────────────────────────────────────
    h1(doc, "7.  The three seeds")
    table(doc, ["seed", "decides", "changing it"], [
        ("seed", "which devices are drawn",
         "a different pool, and a different pool folder"),
        ("split_seed", "which devices are held out",
         "a different assignment of the SAME devices; it does not redraw "
         "them"),
        ("train_seed", "initial weights and batch order",
         "the same data, a different place for the optimisation to land"),
    ], widths=[1.0, 1.9, 3.5])
    para(doc,
         "train_seed is the one to be careful with. One run per arm is one "
         "roll of the dice, and a gap between two arms cannot be told apart "
         "from the spread of that roll. model/ and evaluation/ carry the "
         "seed in their folder names, so an arm can be trained at several "
         "seeds and all the results survive — compare the clusters, not the "
         "single numbers. The validation slice is deliberately pinned to a "
         "fixed seed, so every repeat of an arm is scored against the same "
         "validation devices.")

    # 8 ────────────────────────────────────────────────────────────────────
    h1(doc, "8.  What ends up on disk")
    code(doc,
         ("data/_device_pools/%s/   shared, expensive\n" % pool_example()) +
         "results/<sweep>/\n"
         "    <budget>/                     one folder per configuration\n"
         "        config.json               the settings, read back by 2 and 3\n"
         "        train.npz  test.npz       the measured inputs and the answers\n"
         "        dataset_summary.json      split and separation evidence\n"
         "        figures/                  what the dataset looks like\n"
         "        model/unet.pt             weights + the chosen threshold\n"
         "        evaluation/               metrics.json, per_device.csv, figures\n"
         "    comparison.csv                one row per configuration\n"
         "    figures/<bundle>/             the gallery\n"
         "    model_structure.yaml          what the network is\n"
         "    hyperparameters.yaml          what it was run with")
    para(doc,
         "One run is one folder. The device pools are deliberately left "
         "outside it: they are the expensive part and they do not depend on "
         "the budget.")

    # 9 ────────────────────────────────────────────────────────────────────
    h1(doc, "9.  Running it, and checking it")
    code(doc,
         "python scripts/run_0_full_sweep.py      steps 1-4 for a whole grid\n"
         "python scripts/run_1_generate_dataset.py   one budget's dataset\n"
         "python scripts/run_2_train_model.py        train\n"
         "python scripts/run_3_evaluate_model.py     score (compares after)\n"
         "python scripts/run_9_update_readme.py      README results, from the run\n"
         "pytest -q                                  %d tests, no data needed"
         % test_count())
    para(doc,
         "run_0 takes every setting it does not own from run_1's CONFIG and "
         "calls the same library functions in the same order, so it cannot "
         "produce a different answer than running the stages by hand. Check "
         "its FRESH_START and WIPE_DEVICE_POOLS flags before running: "
         "together they empty results/ and data/.")
    para(doc,
         "After any study run, re-run run_9 — otherwise the README's "
         "results section describes a run that no longer exists.")

    # 10 ───────────────────────────────────────────────────────────────────
    h1(doc, "10.  Appendix — the public surface, module by module")
    para(doc,
         "Generated by walking the source: every top-level class and "
         "function whose name does not start with an underscore.")
    for rel, _path, summary, lines, public in mods:
        if not public:
            continue
        h2(doc, rel.replace("csdrecon/", ""))
        if summary:
            p = para(doc, summary)
            p.runs[0].font.color.rgb = GREY
            p.paragraph_format.space_after = Pt(2)
        code(doc, "  ".join(public))

    doc.save(OUT_DOCX)
    return OUT_DOCX, total, len(mods)


# ── PDF, through Word ─────────────────────────────────────────────────────

PS = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
foreach ($src in $args) {
  $pdf = [System.IO.Path]::ChangeExtension($src, '.pdf')
  try {
    $doc = $word.Documents.Open($src, $false, $true)
    $n = $doc.ComputeStatistics(2)
    $doc.ExportAsFixedFormat($pdf, 17)
    $doc.Close(0)
    Write-Output ("OK`t{0}`t{1}" -f [System.IO.Path]::GetFileName($src), $n)
  } catch {
    Write-Output ("SKIP`t{0}`t{1}" -f [System.IO.Path]::GetFileName($src), $_.Exception.Message)
  }
}
$word.Quit()
"""


def to_pdf(docx_path):
    script = os.path.join(HERE, "_export_pdf.ps1")
    with open(script, "w", encoding="utf-8") as fh:
        fh.write(PS)
    out = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
         "-File", script, docx_path],
        capture_output=True, text=True)
    os.remove(script)
    for line in out.stdout.splitlines():
        if line.startswith("OK"):
            _, name, pages = (line.split("\t") + ["", ""])[:3]
            return os.path.splitext(docx_path)[0] + ".pdf", pages
        if line.startswith("SKIP"):
            print("  not exported:", line.split("\t")[-1][:120])
    if out.stderr.strip():
        print(out.stderr.strip()[:400])
    return None, None


def main():
    print("The codebase, explained")
    print("=" * 74)
    docx_path, lines, n_mods = build()
    print("  %-34s %d modules, %d lines described"
          % (os.path.basename(docx_path), n_mods, lines))
    pdf, pages = to_pdf(docx_path)
    if pdf:
        print("  %-34s %s pages" % (os.path.basename(pdf), pages))
    print()
    print("->", HERE)


if __name__ == "__main__":
    main()
